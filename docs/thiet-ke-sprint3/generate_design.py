"""Generate the S3-L1 design artifacts from one data dictionary; no database writes."""
from pathlib import Path
import json, subprocess, html
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2]
DOCS=Path(__file__).parent
OUT=ROOT/'Sản phẩm yêu cầu/Lê Hoàng Lâm/S3-L1'
OUT.mkdir(parents=True,exist_ok=True)
# name, SQL type, NULL?, key/default/rules
T={
'TaiKhoan':[
('MaTaiKhoan','INT','Không','PK, IDENTITY(1,1)'),
('Email','NVARCHAR(100)','Không','UNIQUE; lưu email đã Trim/chuẩn hóa thống nhất với đăng nhập'),
('MatKhauHash','NVARCHAR(255)','Không','Chỉ lưu chuỗi hash của PasswordHasher; không có cột mật khẩu rõ'),
('VaiTro','TINYINT','Không','Giữ mã 1–4 của nguồn hiện có; đối chiếu lại S3-H1 trước triển khai'),
('TrangThai','TINYINT','Không','DEFAULT 1; CHECK IN (0,1): khóa/hoạt động'),
('NgayTao','DATETIME2','Không','DEFAULT SYSDATETIME()')],
'GioHang':[
('MaKhachHang','INT','Không','PK ghép, FK → KhachHang.MaKhachHang'),
('MaSach','INT','Không','PK ghép, FK → Sach.MaSach'),
('SoLuong','INT','Không','CHECK > 0; khi mua phải kiểm tra lại hàng/giới hạn nghiệp vụ'),
('NgayThem','DATETIME2','Không','DEFAULT SYSDATETIME()')],
'DonHang':[
('MaDon','INT','Không','PK, IDENTITY(1,1); giữ tên khóa của nguồn hiện có'),
('MaKhachHang','INT','Không','FK → KhachHang.MaKhachHang'),
('NgayDat','DATETIME2','Không','DEFAULT SYSDATETIME(); thời điểm chốt đơn'),
('TrangThai','TINYINT','Không','DEFAULT 1; CHECK 1–5 theo bảng trạng thái hiện có'),
('TenNguoiNhan','NVARCHAR(100)','Không','Bản sao tên người nhận tại lúc đặt'),
('SoDienThoaiNhan','VARCHAR(15)','Không','Bản sao số điện thoại, kiểm tra định dạng ở ứng dụng'),
('DiaChiNhan','NVARCHAR(255)','Không','Số nhà/đường/địa chỉ chi tiết, bản sao tại lúc đặt'),
('TinhThanh','NVARCHAR(100)','Không','Bản sao tên Tỉnh/Thành đã kiểm tra; đề xuất tăng từ 50 lên 100 ký tự'),
('PhuongXa','NVARCHAR(100)','Không','BỔ SUNG: bản sao Phường/Xã tại lúc đặt; bắt buộc với đơn mới'),
('PhiGiao','DECIMAL(12,0)','Không','CHECK >= 0; PHÍ CHỐT, không tính lại từ biểu phí khi đọc đơn'),
('TuyenGiao','TINYINT','Không','BỔ SUNG: FK → BangPhiGiao.Tuyen; tuyến đã chọn lúc chốt'),
('KhoiLuongTinhPhiGram','INT','Không','BỔ SUNG: CHECK > 0; khối lượng tính phí đã chốt'),
('LyDoHuy','NVARCHAR(255)','Có','Giữ theo nguồn; bắt buộc theo quy trình khi hủy nếu đặc tả yêu cầu'),
('NgayCapNhat','DATETIME2','Không','DEFAULT SYSDATETIME()'),
('NguoiCapNhat','INT','Có','FK → NhanVien.MaNhanVien; bảng nội bộ có sẵn ngoài phạm vi')],
'ChiTietDonHang':[
('MaDon','INT','Không','PK ghép, FK → DonHang.MaDon'),
('MaSach','INT','Không','PK ghép, FK → Sach.MaSach'),
('SoLuong','INT','Không','CHECK > 0; số lượng đã đặt'),
('DonGia','DECIMAL(12,0)','Không','CHECK > 0; GIÁ LÚC ĐẶT, không đọc Sach.GiaBan để tính lại đơn cũ')],
'DiaChiKhachHang':[
('MaDiaChi','INT','Không','PK, IDENTITY(1,1)'),
('MaKhachHang','INT','Không','FK → KhachHang.MaKhachHang; kiểm tra chủ sở hữu trước đọc/sửa'),
('TenNguoiNhan','NVARCHAR(100)','Không','Thông tin địa chỉ đang lưu, có thể được khách chỉnh sửa'),
('SoDienThoai','VARCHAR(15)','Không','Kiểm tra định dạng ở ứng dụng'),
('DiaChiChiTiet','NVARCHAR(255)','Không','Số nhà/đường/chi tiết hiện tại'),
('TinhThanh','NVARCHAR(100)','Không','Tên Tỉnh/Thành hiện tại; kiểm tra theo danh mục hành chính thật'),
('PhuongXa','NVARCHAR(100)','Không','BỔ SUNG: Phường/Xã thuộc Tỉnh/Thành đã chọn'),
('MacDinh','BIT','Không','DEFAULT 0; unique index lọc theo MaKhachHang WHERE MacDinh=1')],
'BangPhiGiao':[
('Tuyen','TINYINT','Không','PK; CHECK 1–4: nội tỉnh/nội miền/đặc biệt/liên miền theo nguồn'),
('TenTuyen','NVARCHAR(50)','Không','Tên tuyến; không dùng tên làm khóa'),
('PhiDen1000g','DECIMAL(12,0)','Không','CHECK >= 0'),
('PhiDen1500g','DECIMAL(12,0)','Không','CHECK >= 0'),
('PhiDen2000g','DECIMAL(12,0)','Không','CHECK >= 0'),
('PhiMoi500gTiep','DECIMAL(12,0)','Không','CHECK >= 0; mỗi 500g vượt 2000g, làm tròn lên')],
'CauHinhCuaHang':[
('MaCauHinh','TINYINT','Không','PK; CHECK MaCauHinh=1; cấu hình 1 cửa hàng'),
('TenCuaHang','NVARCHAR(150)','Không','Tên hiển thị của cửa hàng'),
('DiaChiKho','NVARCHAR(255)','Không','Địa chỉ gửi chi tiết đang áp dụng'),
('TinhThanhKho','NVARCHAR(100)','Không','Điểm gửi để xác định tuyến; kiểm tra danh mục hành chính'),
('PhuongXaKho','NVARCHAR(100)','Không','Phường/Xã điểm gửi'),
('KhoiLuongDongGoiGram','INT','Không','DEFAULT 0; CHECK >=0; phụ phí trọng lượng bao bì theo cấu hình'),
('NgayCapNhat','DATETIME2','Không','DEFAULT SYSDATETIME()')]
}
FK=[
('GioHang','MaKhachHang','KhachHang','MaKhachHang'),('GioHang','MaSach','Sach','MaSach'),
('DonHang','MaKhachHang','KhachHang','MaKhachHang'),('DonHang','TuyenGiao','BangPhiGiao','Tuyen'),
('DonHang','NguoiCapNhat','NhanVien','MaNhanVien'),('ChiTietDonHang','MaDon','DonHang','MaDon'),
('ChiTietDonHang','MaSach','Sach','MaSach'),('DiaChiKhachHang','MaKhachHang','KhachHang','MaKhachHang')]
EXT={'KhachHang':{'MaKhachHang':'INT','MaTaiKhoan':'INT'},'Sach':{'MaSach':'INT'},'NhanVien':{'MaNhanVien':'INT'}}
# Consistency check for references in the proposed dictionary; not a SQL Server execution test.
for child,col,parent,pk in FK:
 fields={x[0]:x[1] for x in T[child]}
 target={x[0]:x[1] for x in T[parent]} if parent in T else EXT[parent]
 assert fields[col]==target[pk],(child,col,parent,pk)
assert EXT['KhachHang']['MaTaiKhoan']==dict((x[0],x[1]) for x in T['TaiKhoan'])['MaTaiKhoan']
(DOCS/'schema-de-xuat.json').write_text(json.dumps({'status':'Đề xuất thiết kế; chưa triển khai migration','tables':T,'foreign_keys':FK,'physical_name_mapping':{'ChiTietDonHang':'ChiTietDon trong nguồn hiện có'}},ensure_ascii=False,indent=2))

def graph(detail=False):
 lines=['digraph G {','graph [rankdir=LR, bgcolor="white", pad="0.25", nodesep="0.35", ranksep="0.8", fontname="DejaVu Sans"];','node [shape=plain,fontname="DejaVu Sans"];','edge [fontname="DejaVu Sans",fontsize=10,color="#425466",arrowsize=0.7];']
 allnodes=list(T)+['KhachHang','Sach','NhanVien']
 if not detail:allnodes+=['TheLoai','NhaXuatBan','TacGia','SachTacGia']
 for name in allnodes:
  color='#dbeafe' if name in T else '#e5e7eb'
  if name in T:
   if detail:attrs=[f'{a} : {b}'+(' [PK]' if 'PK' in k else '') for a,b,n,k in T[name]]
   else:
    show={'TaiKhoan':['MaTaiKhoan','Email','MatKhauHash'],'DonHang':['MaDon','MaKhachHang','PhuongXa','TinhThanh','PhiGiao'],'ChiTietDonHang':['MaDon','MaSach','SoLuong','DonGia'],'GioHang':['MaKhachHang','MaSach','SoLuong'],'DiaChiKhachHang':['MaDiaChi','MaKhachHang','PhuongXa','TinhThanh'],'BangPhiGiao':['Tuyen','PhiDen1000g','PhiMoi500gTiep'],'CauHinhCuaHang':['MaCauHinh','TinhThanhKho','PhuongXaKho']}
    attrs=[a for a,b,n,k in T[name] if a in show[name]]
  else:
   attrs={'KhachHang':['PK MaKhachHang','FK UNIQUE MaTaiKhoan'],'Sach':['PK MaSach','FK MaTheLoai / MaNXB','GiaBan / KhoiLuongGram'],'NhanVien':['PK MaNhanVien'],'TheLoai':['PK MaTheLoai'],'NhaXuatBan':['PK MaNXB'],'TacGia':['PK MaTacGia'],'SachTacGia':['PK/FK MaSach, MaTacGia']}[name]
  title=name+(' (*)' if name=='ChiTietDonHang' else '')
  rows=''.join('<TR><TD ALIGN="LEFT">'+html.escape(a)+'</TD></TR>' for a in attrs)
  lines.append(f'{name} [label=<<TABLE BORDER="1" CELLBORDER="0" CELLSPACING="0" CELLPADDING="5"><TR><TD BGCOLOR="{color}"><B>{title}</B></TD></TR>{rows}</TABLE>>];')
 # Edges point from referenced parent to dependent child; labels give exact cardinality.
 lines.append('TaiKhoan -> KhachHang [label="1 : 0..1"];')
 for child,col,parent,pk in FK:
  label=f'1 : 0..n ({col})'
  if child=='DonHang' and col=='NguoiCapNhat':label='0..1 nhân viên / đơn; 0..n đơn / nhân viên'
  if child=='ChiTietDonHang' and col=='MaDon':label='1 : 1..n (nghiệp vụ; MaDon)'
  lines.append(f'{parent} -> {child} [label="{label}"];')
 if not detail:
  lines+=['TheLoai -> Sach [label="1 : 0..n"];','NhaXuatBan -> Sach [label="1 : 0..n"];','Sach -> SachTacGia [label="1 : 0..n"];','TacGia -> SachTacGia [label="1 : 0..n"];']
 lines+=['DiaChiKhachHang -> DonHang [style=dashed,color="#aa6000",label="sao chép địa chỉ lúc đặt\n(không phải FK)",constraint=false];','CauHinhCuaHang -> DonHang [style=dashed,color="#aa6000",label="điểm gửi / đóng gói để tính phí\n(không phải FK)",constraint=false];','}']
 return '\n'.join(lines)
for name,detail in [('erd-noi-voi-quan-ly-sach',False),('erd-chi-tiet-sprint3',True)]:
 p=DOCS/(name+'.dot');p.write_text(graph(detail))
 for ext in ['svg','png']:
  subprocess.run(['dot','-T'+ext,'-Gdpi=160',str(p),'-o',str(DOCS/(name+'.'+ext))],check=True)

# Shared content events used to generate the Word and Markdown versions.
B=[]
def h(t,l=1):B.append(('h',l,t))
def p(t):B.append(('p',t))
def table(headers,rows):B.append(('table',headers,rows))
def code(t):B.append(('code',t))
def image(name,caption,notes):B.append(('image',name,caption,notes))
h('S3-L1 – Thiết kế cơ sở dữ liệu Sprint 3 theo 4 bước')
p('Người phụ trách: Lê Hoàng Lâm. Ngày lập: 21/10/2026 (Asia/Saigon). Phạm vi: 7 bảng theo thẻ, nối với KhachHang, Sach và các bảng quản lý sách hiện có.')
h('0. Nguồn, phạm vi và trạng thái thiết kế',2)
p('Đây là thiết kế đề xuất có đối chiếu nguồn đang có, không phải xác nhận đã triển khai hoặc đã Pass. Hồ sơ chưa có WebsiteBanSach_Sprint3/database/01_tao_bang.sql, docs/dac_ta_sprint3_gio_hang_dat_hang.md mục 4, tài liệu S4-L1 riêng và đặc tả S3-H1/S3-Y1 cuối. Vì vậy, các phần bổ sung được đánh dấu để nhóm đối chiếu trước khi viết migration và code.')
p('Nguồn đã đọc: database/01_tao_bang.sql của master trong WebsiteBanSach-S4-L4.zip, báo cáo chung và mẫu cá nhân đã cung cấp. SQL nền có tên CSDL NhaSachOnline_S3, nhưng tên này không chứng minh đó là đúng bản tham khảo WebsiteBanSach_Sprint3 nêu trong thẻ. SHA-256 SQL nền: 5cb74aa1e7441b4653279ab3d2c07dea78bcb1a77fa1ff2c4491ec08d226f601.')
table(['Điểm đối chiếu','Nguồn hiện có','Thiết kế đề xuất'],[
('Chi tiết đơn','Tên vật lý ChiTietDon, khóa (MaDon,MaSach)','Tên logic theo thẻ: ChiTietDonHang. Khi triển khai giữ bảng cũ qua ánh xạ hoặc thống nhất migration đổi tên; không tạo đồng thời hai bảng.'),
('Địa chỉ','DiaChiKhachHang và DonHang chỉ có TinhThanh','Bổ sung PhuongXa riêng ở cả địa chỉ hiện tại và bản sao địa chỉ trên đơn.'),
('Cấu hình cửa hàng','Chưa có CauHinhCuaHang','Đề xuất 1 dòng cấu hình điểm gửi và khối lượng đóng gói.'),
('Giá và phí lịch sử','ChiTietDon.DonGia, DonHang.PhiGiao đã có','Giữ là giá/phí lúc đặt, không tính lại từ Sach/BangPhiGiao hiện hành.'),
('Giỏ khách chưa đăng nhập','Bảng GioHang gắn MaKhachHang NOT NULL','Giỏ tạm ở Session/cookie theo UC03; sau đăng nhập mới hợp nhất vào GioHang; không cho FK khách NULL.'),
('Tuyến và trọng lượng chốt','DonHang chưa có hai trường này','Đề xuất TuyenGiao và KhoiLuongTinhPhiGram để giải thích phí đã chốt.')])
h('1. Bước 1 – Xác định thực thể, khóa và thuộc tính')
table(['Bảng','Vì sao tồn tại','Khóa và dữ liệu chính'],[
('TaiKhoan','Tách danh tính đăng nhập khỏi địa chỉ, giỏ và đơn; dùng cho UC15/UC26.','MaTaiKhoan; Email duy nhất; MatKhauHash; vai trò; trạng thái.'),
('GioHang','Lưu lựa chọn mua hiện tại của khách, chưa phải cam kết giá/đơn hàng. Một dòng là một sách trong giỏ.','PK (MaKhachHang, MaSach); SoLuong; NgayThem.'),
('DonHang','Lưu giao dịch và bản sao người nhận/phí ở thời điểm chốt; không phụ thuộc địa chỉ khách sửa về sau.','MaDon; khách; thời điểm; trạng thái; người nhận; địa chỉ; PhiGiao.'),
('ChiTietDonHang','Biểu diễn n–m giữa đơn và sách, giữ số lượng và giá lúc đặt cho từng dòng.','PK (MaDon,MaSach); SoLuong; DonGia.'),
('DiaChiKhachHang','Khách có nhiều địa chỉ và một địa chỉ mặc định tùy chọn; có thể sửa/xóa mà không đổi đơn cũ.','MaDiaChi; khách sở hữu; tên/SĐT; chi tiết; Tỉnh/Thành; Phường/Xã; MacDinh.'),
('BangPhiGiao','Quản lý biểu phí hiện hành theo tuyến và ngưỡng cân nặng, tránh hard-code giá trong Controller.','Tuyen; tên tuyến; các mức phí đến 1000g/1500g/2000g và mỗi 500g tiếp.'),
('CauHinhCuaHang','Giữ thông tin một cửa hàng/điểm gửi, làm đầu vào xác định tuyến và cân nặng đóng gói.','MaCauHinh=1; tên; địa chỉ kho; Tỉnh/Thành, Phường/Xã; trọng lượng đóng gói.')])
p('KhachHang là bảng có sẵn bắt buộc làm cầu nối từ TaiKhoan tới giỏ, địa chỉ và đơn. Không gán MaTaiKhoan trực tiếp vào FK MaKhachHang. NhanVien là bảng có sẵn để ghi người cập nhật đơn. Các bảng này được vẽ là phụ thuộc ngoài phạm vi, không thiết kế lại thành bảng mới.')
p('Một dòng Sach vẫn là một ấn bản; MaSach trong giỏ và chi tiết đơn cùng tham chiếu PK Sach.MaSach. Không nối trực tiếp tới TacGia hoặc một CuonSach vật lý trong mô-đun đặt hàng; gán cuốn vật lý khi đóng gói là phần kho hiện có.')
h('2. Bước 2 – Xác định quan hệ và lực lượng tham gia')
table(['Quan hệ','Bội số và FK','Quy tắc'],[
('TaiKhoan – KhachHang','1 : 0..1; KhachHang.MaTaiKhoan FK UNIQUE','Tài khoản nội bộ có thể không có hồ sơ khách; mỗi KhachHang gắn đúng một tài khoản.'),
('KhachHang – DiaChiKhachHang','1 : 0..n; MaKhachHang FK','Một địa chỉ thuộc đúng một khách; ứng dụng phải kiểm tra quyền sở hữu.'),
('KhachHang – GioHang – Sach','Hai quan hệ 1 : 0..n; PK ghép của GioHang','Một khách không có hai dòng giỏ cùng MaSach; giỏ trống là hợp lệ.'),
('KhachHang – DonHang','1 : 0..n; DonHang.MaKhachHang FK','Mỗi đơn thuộc đúng một khách; đơn không chuyển sang khách khác.'),
('DonHang – ChiTietDonHang','1 : 1..n về nghiệp vụ; FK chỉ đảm bảo mỗi chi tiết có một đơn','Đơn chốt phải có ít nhất một dòng. Ghi cả đơn/chi tiết trong một giao dịch; FK một mình không ngăn đơn rỗng.'),
('Sach – ChiTietDonHang','1 : 0..n; ChiTietDonHang.MaSach FK','Sách từng có trong đơn không được xóa cứng; dùng trạng thái ngừng kinh doanh.'),
('BangPhiGiao – DonHang','1 : 0..n; TuyenGiao FK (đề xuất)','Lưu tuyến để tra cứu; PhiGiao là bản sao độc lập với các mức phí thay đổi sau này.'),
('NhanVien – DonHang','1 : 0..n; NguoiCapNhat FK NULL được','Có thể chưa có người nội bộ xử lý khi khách vừa đặt.'),
('DiaChiKhachHang → DonHang','Luồng sao chép, KHÔNG là FK bắt buộc','Chọn địa chỉ hiện tại rồi sao chép đầy đủ; có thể nhập địa chỉ nhận riêng.'),
('CauHinhCuaHang → chốt đơn','Luồng đọc cấu hình, KHÔNG là FK','CHECK MaCauHinh=1 chỉ đảm bảo tối đa một dòng; khởi tạo phải bảo đảm có một dòng hợp lệ trước đặt hàng.')])
p('Các FK của đơn/chi tiết dùng NO ACTION khi xóa: không xóa dây chuyền lịch sử từ tài khoản hoặc sách. Địa chỉ hiện tại không bị đơn cũ giữ FK, nên khách có thể xóa một địa chỉ sau khi ứng dụng kiểm tra quyền sở hữu; bản sao trên đơn không đổi.')
h('2.1. ERD tổng quan nối với phần quản lý sách',2)
image('erd-noi-voi-quan-ly-sach.png','Hình 1: ERD Sprint 3 nối với Sach và nhóm danh mục/tác giả',[
('Mục đích','Thể hiện điểm nối tài khoản, khách, giỏ, đơn và sách; không tạo một bảng sản phẩm khác.'),
('Trạng thái ban đầu','Có schema quản lý sách và một phần tài khoản/đơn trong SQL nền; chưa có schema Sprint 3 cuối.'),
('Thao tác và dữ liệu nhập','Đối chiếu PK/FK, giữ cầu nối KhachHang, thêm PhuongXa và cấu hình theo đề xuất.'),
('Kết quả nhìn thấy','GioHang và ChiTietDonHang cùng nối Sach; các bảng sách hiện có màu xám, 7 bảng phạm vi màu xanh. Mũi tên liền đi từ bảng cha tới bảng con.'),
('Nhận xét','Nhãn 1:0..n nói một cha có thể chưa có dòng con; mỗi dòng con có đúng một cha khi FK NOT NULL. NguoiCapNhat có thể NULL; tối thiểu 1 chi tiết/đơn là quy tắc nghiệp vụ cần giao dịch, không tự được FK đảm bảo. Nét đứt là sao chép/đọc dữ liệu, không phải FK.')])
h('3. Bước 3 – Chuẩn hóa 1NF, 2NF, 3NF')
table(['Mức','Áp dụng','Giải thích'],[
('1NF','Không lưu danh sách MaSach trong DonHang; không ghép nhiều địa chỉ vào TaiKhoan.','Mỗi dòng ChiTietDonHang/GioHang là một sách. Mỗi địa chỉ là một dòng DiaChiKhachHang. TinhThanh và PhuongXa là hai thuộc tính riêng.'),
('2NF','GioHang có khóa (MaKhachHang,MaSach); ChiTietDonHang có khóa (MaDon,MaSach).','Số lượng giỏ phụ thuộc cả khách lẫn sách. SoLuong/DonGia của chi tiết phụ thuộc cả đơn lẫn sách, không chỉ MaSach.'),
('3NF','Tên sách, tác giả, thể loại, NXB giữ trong nhóm sách; tài khoản không chứa dữ liệu đơn.','Không lặp các tên danh mục trong giỏ/chi tiết. BangPhiGiao và CauHinhCuaHang tách khỏi giao dịch; thông tin hiện hành không bị nhân bản trên mọi giỏ.'),
('Tổng tiền tính được','Không cần lưu TongTienSach/TongThanhToan trong DonHang.','Tổng = SUM(SoLuong × DonGia lịch sử) + PhiGiao đã chốt. Đọc đơn không JOIN Sach.GiaBan để tính tiền.'),
('Bản sao có chủ đích','Giữ DonGia, PhiGiao và địa chỉ nhận ngay trên giao dịch.','Đây là sự kiện tại thời điểm đặt, không phải bản sao để đồng bộ theo dữ liệu hiện tại. Thay đổi sách/biểu phí/địa chỉ không được cập nhật ngược đơn cũ.')])
p('Phụ thuộc hàm lịch sử: (MaDon,MaSach) → DonGia; MaDon → TenNguoiNhan, SoDienThoaiNhan, DiaChiNhan, TinhThanh, PhuongXa, PhiGiao. Các giá trị này KHÔNG phụ thuộc vào giá/địa chỉ hiện hành. Vì vậy không loại bỏ snapshot dưới danh nghĩa chuẩn hóa.')
p('MaDiaChi → MaKhachHang và thông tin địa chỉ hiện tại; MaTaiKhoan → Email/MatKhauHash; Tuyen → các mức phí hiện hành; MaCauHinh=1 → thông tin cửa hàng. Nếu bổ sung danh mục hành chính có mã thật, thiết kế FK tỉnh/xã riêng; bộ mã và quan hệ xã–tỉnh chưa được cung cấp nên không tự bịa mã trong thiết kế này.')
h('4. Bước 4 – ERD chi tiết và thiết kế từng bảng')
p('Các kiểu dữ liệu dưới đây là thiết kế đích trên SQL Server cho triển khai mới. Không chạy migration trên dữ liệu nhóm trong phiên này. BỔ SUNG đánh dấu trường chưa có trong SQL nền. NVARCHAR dùng cho tiếng Việt; tiền VND dùng DECIMAL, không dùng FLOAT/REAL.')
for i,(name,rows) in enumerate(T.items(),1):
 h(f'4.{i}. {name}',2)
 if name=='ChiTietDonHang':p('Tên logic theo thẻ; bảng vật lý nguồn là ChiTietDon. Các kiểu/khóa bên dưới áp dụng cho cùng một thực thể, không đề xuất duy trì hai bảng song song.')
 if name=='CauHinhCuaHang':p('Bảng đề xuất mới, chưa được khẳng định là cấu trúc của bản tham khảo Sprint 3. Thiết kế cho một cửa hàng; mở rộng nhiều kho cần mô hình riêng.')
 table(['Cột','Kiểu SQL','NULL','Khóa, mặc định và ràng buộc'],rows)
h('4.8. Chỉ mục và ràng buộc bổ sung',2)
code('''-- Tối đa một địa chỉ mặc định mỗi khách; không bắt buộc khách phải có địa chỉ.
CREATE UNIQUE INDEX UX_DiaChiKhachHang_MacDinh
ON DiaChiKhachHang(MaKhachHang) WHERE MacDinh = 1;

-- Tra lịch sử đặt hàng theo khách/thời gian.
CREATE INDEX IX_DonHang_Khach_NgayDat
ON DonHang(MaKhachHang, NgayDat DESC);

-- Nếu tên vật lý được thống nhất là ChiTietDonHang:
CREATE INDEX IX_ChiTietDonHang_MaSach
ON ChiTietDonHang(MaSach);
''')
p('Trích đoạn là minh họa thiết kế, chưa thực thi, không thay thế migration. PK/UNIQUE đã tạo chỉ mục cho email, cặp khách–sách và cặp đơn–sách. Các mức phí cần CHECK >=0; các số lượng CHECK >0. Đổi địa chỉ mặc định phải bỏ cờ cũ và đặt cờ mới trong cùng giao dịch để không vi phạm unique index.')
p('Kiểm tra TinhThanh–PhuongXa phải dựa trên danh mục hành chính thật được nhóm chọn; độ dài chuỗi và NOT NULL không chứng minh hai tên thuộc nhau. Nếu chỉ dùng text như phạm vi hiện tại, ứng dụng thực hiện kiểm tra cặp tên từ nguồn chuẩn trước khi lưu.')
h('4.9. ERD chi tiết bảy bảng',2)
image('erd-chi-tiet-sprint3.png','Hình 2: Thuộc tính và khóa của mô-đun tài khoản–giỏ–đơn',[
('Mục đích','Đối chiếu từ điển dữ liệu với khóa và các liên kết của 7 bảng.'),
('Trạng thái ban đầu','Đã chốt nội bộ thiết kế đề xuất trong tài liệu, còn chờ đặc tả S3-H1/S3-Y1 để review.'),
('Thao tác và dữ liệu nhập','Sinh sơ đồ từ cùng từ điển dữ liệu JSON; đối chiếu kiểu PK/FK và tên trường.'),
('Kết quả nhìn thấy','MatKhauHash, DonGia, PhiGiao, TinhThanh, PhuongXa đều hiện trong các bảng tương ứng. KhachHang/Sach/NhanVien là bảng đã có.'),
('Nhận xét','Các ràng buộc CHECK, UNIQUE, NULL xem chi tiết bảng; sơ đồ lược bỏ một số nhãn để dễ đọc. Có SVG và DOT kèm theo để phóng to hoặc sửa. Đây không phải hình chụp SQL Server đang chạy.')])
h('5. Quy tắc chốt đơn và giữ dữ liệu lịch sử')
h('5.1. Mật khẩu chỉ lưu bản băm',2)
p('UC15: sau kiểm tra email hợp lệ/duy nhất, dùng PasswordHasher của ASP.NET Core để tạo chuỗi hash chứa salt và tham số thuật toán rồi lưu MatKhauHash. UC26: VerifyHashedPassword đối chiếu mật khẩu nhập với hash; xử lý nâng cấp hash theo API nếu cần. Không lưu mật khẩu rõ, không lưu mã hóa có thể giải ngược, không dùng SHA-256 thuần. Không ghi mật khẩu hoặc hash thật vào tài liệu, log và dữ liệu minh họa. Cơ chế khóa đăng nhập/chính sách mật khẩu chi tiết phải khớp S3-H1; chưa tự thêm quy tắc giả định 10 lần/15 phút.')
h('5.2. Giỏ hàng và giá lúc đặt',2)
p('GioHang chỉ giữ MaSach/SoLuong, không giữ một giá được coi là cam kết. Khi mở giỏ có thể đọc giá hiện hành để hiển thị; khi bấm đặt phải đọc/kiểm tra lại giá ở máy chủ. Nếu giá/khả dụng đổi, báo khách xem lại trước xác nhận, không tin DonGia hay PhiGiao do client gửi. Với khách chưa đăng nhập, giỏ tạm được hợp nhất theo MaSach vào giỏ tài khoản; giới hạn theo số lượng hợp lệ và tồn khả dụng, không tạo hai dòng trùng khóa.')
p('ChiTietDonHang.DonGia là giá đã được chấp nhận và chốt lúc tạo đơn. Một sách giá 99.000đ lúc đặt sau đó đổi 120.000đ: đơn cũ vẫn tính 99.000đ cho dòng đó. Việc sách ngừng kinh doanh không làm mất chi tiết đơn hoặc đổi giá lịch sử.')
h('5.3. Phí giao và địa chỉ nhận',2)
p('Trọng lượng tính phí w = SUM(SoLuong × Sach.KhoiLuongGram tại lúc chốt) + CauHinhCuaHang.KhoiLuongDongGoiGram. Tất cả phép tính ở server, dùng số nguyên gram và DECIMAL tiền. w phải >0. Xác định tuyến từ điểm gửi và Tỉnh/Thành nhận theo bảng quy tắc S3-Y1; không đoán tuyến từ thứ tự tên tỉnh hoặc tự đặt khoảng cách. Các tuyến đặc biệt và bộ dữ liệu tỉnh/xã cần nguồn nhóm xác nhận.')
code('''Nếu w <= 1000: phi = PhiDen1000g
Nếu 1000 < w <= 1500: phi = PhiDen1500g
Nếu 1500 < w <= 2000: phi = PhiDen2000g
Nếu w > 2000:
    phi = PhiDen2000g + ceil((w - 2000) / 500) * PhiMoi500gTiep
''')
p('Đây là cách đọc biểu phí theo ngưỡng trong SQL nguồn, cần đối chiếu mục 4 của đặc tả Sprint 3 trước khi lập trình. Lưu kết quả vào DonHang.PhiGiao cùng TuyenGiao/KhoiLuongTinhPhiGram. Không đọc lại phí hiện hành để sửa tổng đơn cũ. Không tự thêm miễn phí giao, thuế hoặc khuyến mại khi chưa có đặc tả.')
p('Khi chốt, sao chép TenNguoiNhan, SoDienThoaiNhan, DiaChiNhan, TinhThanh, PhuongXa từ lựa chọn đã kiểm tra vào DonHang. Đơn không chỉ lưu MaDiaChi rồi JOIN địa chỉ hiện tại khi hiển thị. Địa chỉ có thể được nhập riêng trong checkout nếu UC04 cho phép; không bắt buộc phải lưu trước vào sổ địa chỉ. Sửa địa chỉ một đơn theo quy trình nghiệp vụ là thao tác riêng có phân quyền/lịch sử, không phải đồng bộ tự động khi khách sửa sổ địa chỉ.')
h('5.4. Một giao dịch đặt hàng nguyên tử',2)
table(['Thứ tự','Xử lý trong giao dịch'],[
('1','Xác định MaKhachHang từ tài khoản đã xác thực, kiểm tra các dòng được chọn thuộc giỏ của khách; giỏ được chọn không rỗng.'),
('2','Đọc lại sách đang bán, giá, trọng lượng và số lượng khả dụng; kiểm tra phiên bản/điều kiện cập nhật để tránh đọc giá cũ hoặc bán vượt tồn.'),
('3','Kiểm tra địa chỉ, cặp Tỉnh/Thành–Phường/Xã, quyền sở hữu địa chỉ; đọc cấu hình/biểu phí hợp lệ và tính phí ở server.'),
('4','Tạo DonHang với toàn bộ snapshot và ít nhất một ChiTietDonHang; DonGia lấy từ dữ liệu server đã được xác nhận.'),
('5','Cập nhật/giữ tồn theo đặc tả kho, kiểm tra điều kiện đủ tồn trong câu lệnh cập nhật, không chỉ dựa vào lần đọc trước; xóa đúng các dòng giỏ đã mua.'),
('6','Commit khi tất cả thành công. Có lỗi ở bất kỳ dòng nào: rollback đơn, chi tiết, tồn và thao tác giỏ; không để đơn rỗng, không xóa các sách chưa mua.')])
p('Cơ chế giữ/trừ tồn và chống gửi lại yêu cầu cần đối chiếu UC04/UC05. Bảy bảng này chưa có khóa idempotency; giao dịch một mình không chống được hai POST riêng lẻ giống nhau. Nếu yêu cầu chống tạo hai đơn khi bấm lại/timeout, cần bổ sung mã yêu cầu UNIQUE theo khách hoặc cơ chế đã được nhóm thống nhất; không ghi nhận đã giải quyết trong thiết kế hiện tại.')
h('5.5. Tính tiền khi đọc đơn',2)
code('''-- Tên logic; đổi thành ChiTietDon nếu giữ tên vật lý cũ.
SELECT d.MaDon,
       SUM(CAST(c.SoLuong AS DECIMAL(19,0)) * c.DonGia) AS TongTienSach,
       d.PhiGiao,
       SUM(CAST(c.SoLuong AS DECIMAL(19,0)) * c.DonGia) + d.PhiGiao
           AS TongThanhToan
FROM DonHang d
JOIN ChiTietDonHang c ON c.MaDon = d.MaDon
GROUP BY d.MaDon, d.PhiGiao;
''')
p('Công thức chỉ dùng dòng đơn và phí chốt, không dùng Sach.GiaBan/BangPhiGiao để tính lại lịch sử. Nếu cần lưu thêm tổng tiền để đối soát sau này, phải có lý do nghiệp vụ và một nguồn cập nhật thống nhất; không coi việc lưu tổng tự động là vi phạm 3NF trong mọi trường hợp.')
h('6. Đối chiếu điều kiện đạt và ca kiểm tra thiết kế')
table(['Điều kiện','Vị trí/bằng chứng trong thiết kế'],[
('Giải thích mỗi bảng tồn tại','Bước 1, từ điển 7 bảng và quan hệ Bước 2.'),
('Đơn lưu giá lúc đặt','ChiTietDonHang.DonGia; mục 5.2, công thức 5.5.'),
('Đơn lưu phí giao','DonHang.PhiGiao; không tính lại từ biểu phí hiện hành.'),
('Đơn lưu Tỉnh/Thành, Phường/Xã','Hai trường snapshot trong DonHang, tách với sổ địa chỉ.'),
('Mật khẩu chỉ lưu băm','TaiKhoan.MatKhauHash; quy tắc PasswordHasher tại 5.1.'),
('ERD nối quản lý sách','Hình 1 có GioHang/ChiTietDonHang → Sach → nhóm danh mục/tác giả.'),
('Nhật ký AI','docs/prompt/lam-s3-l1.md ở repo bàn giao.')])
p('Các ca dưới đây là tiêu chí cần chạy khi triển khai, chưa phải kết quả test SQL Server hoặc ứng dụng đã Pass:')
table(['Ca','Thao tác/dữ liệu','Kết quả mong đợi'],[
('TK01','Đăng ký tài khoản, xem dữ liệu lưu','Chỉ có hash mật khẩu; không có mật khẩu rõ trong bảng/log.'),
('TK02','Hai email chỉ khác khoảng trắng/hoa thường theo quy tắc chuẩn hóa','Không tạo hai danh tính trùng; UNIQUE hỗ trợ chặn race.'),
('GH01','Thêm cùng sách hai lần cho cùng khách','Hợp nhất số lượng; không có hai dòng trùng (MaKhachHang,MaSach).'),
('DH01','Giá sách lúc đặt 99.000đ, sau đó đổi 120.000đ','Đơn cũ vẫn có DonGia 99.000đ, tổng không đổi.'),
('DH02','Phí lúc đặt 27.000đ, sau đó biểu phí đổi 30.000đ','Đơn cũ vẫn giữ PhiGiao 27.000đ; số chỉ là minh họa tính bất biến, không phải biểu phí thật.'),
('DC01','Đặt tới Tỉnh A/Phường B, rồi sửa/xóa sổ địa chỉ','Đơn vẫn giữ Tỉnh A/Phường B và người nhận lúc đặt; A/B là dữ liệu giả định.'),
('DC02','Gửi một Phường/Xã không thuộc Tỉnh/Thành đã chọn','Server từ chối trước ghi đơn dựa trên danh mục chuẩn.'),
('DC03','Đặt hai địa chỉ cùng MacDinh=1 cho một khách','Unique index lọc từ chối; giao dịch đổi mặc định giữ nhất quán.'),
('PG01','w=1000,1001,1500,1501,2000,2001,2500,2501 gram','Áp đúng bậc và làm tròn mỗi 500g vượt 2000g; không cộng phí thêm ở đúng mốc.'),
('DH03','Giỏ có hai sách, chỉ chọn một','Đơn có đúng sách được chọn; sách còn lại vẫn ở giỏ.'),
('DH04','Hỏng khi ghi một dòng chi tiết hoặc tồn không đủ','Rollback toàn bộ, không đơn rỗng/không âm tồn/không mất giỏ.'),
('CH01','Không có cấu hình MaCauHinh=1 hoặc biểu phí tuyến','Không chốt đơn với phí tự mặc định 0; báo thiếu cấu hình.'),
('DH05','Hai yêu cầu đồng thời mua số tồn cuối','Chỉ giao dịch có cập nhật đủ điều kiện được chốt; cần kiểm thử cơ chế concurrency thật.')])
h('7. Điểm cần chốt khi tích hợp và bàn giao')
table(['Nội dung','Hướng xử lý trước triển khai'],[
('Nguồn Sprint 3/đặc tả S3-H1, S3-Y1','Đối chiếu schema, chính sách đăng nhập, tuyến phí, danh mục tỉnh/xã, tồn và checkout với bản được nhóm duyệt.'),
('Tên ChiTietDonHang/ChiTietDon','Chọn một tên vật lý; map EF ToTable("ChiTietDon") nếu giữ cũ. Nếu đổi tên phải cập nhật toàn bộ FK, view, SQL, Model và truy vấn trong cùng migration.'),
('Dữ liệu PhuongXa cũ','Triển khai mới dùng NOT NULL. Khi nâng CSDL cũ, thêm NULL tạm và chỉ backfill từ nguồn lịch sử đáng tin; không lấy địa chỉ hiện tại để đoán Phường/Xã của đơn cũ. Chưa đủ dữ liệu thì giữ trạng thái chưa biết cho đơn cũ và chặn đơn mới thiếu phường/xã; chưa ép NOT NULL toàn bảng.'),
('CauHinhCuaHang','Khởi tạo một dòng từ thông tin cửa hàng thật do nhóm cấp; không invent tỉnh/kho/mức phí. CHECK=1 không tự sinh dòng cấu hình.'),
('Phạm vi script hiện có','Không sửa hoặc chạy script DROP TABLE của nhóm. Không có migration hay ứng dụng checkout được triển khai trong nhiệm vụ thiết kế này.'),
('Nguồn GitHub','Sản phẩm thiết kế lưu ở nhánh bàn giao của bom-NO2/chat_gpt đã được cho phép. Đây không phải PR đã merge vào James-Lloyd20/WebsiteBanSach.'),
('Review','Chưa có bình luận review hoặc dấu Pass trên Trello; bảng ca là yêu cầu kiểm tra, không ghi đã chạy.')])
p('Tệp đi kèm: tài liệu Word, Markdown, từ điển JSON, ERD PNG/SVG và mã DOT. Một nguồn từ điển dữ liệu được dùng để sinh bảng và sơ đồ; phần kiểm tra tự động chỉ đối chiếu sự tồn tại/kiểu của PK–FK trong mô hình, không thay cho thực thi SQL Server.')

# Markdown output.
M=[]
for e in B:
 if e[0]=='h': M.append('#'*e[1]+' '+e[2]+'\n')
 elif e[0]=='p': M.append(e[1]+'\n')
 elif e[0]=='code': M.append('```\n'+e[1].rstrip()+'\n```\n')
 elif e[0]=='table':
  head,rows=e[1:];M.append('| '+' | '.join(head)+' |\n| '+' | '.join(['---']*len(head))+' |')
  M.extend('| '+' | '.join(str(x).replace('|','\\|').replace('\n','<br>') for x in row)+' |' for row in rows);M.append('')
 elif e[0]=='image':
  M.append(f'![{e[2]}]({e[1]})\n\n*{e[2]}*\n')
  M.extend('**'+k+':** '+v+'\n' for k,v in e[3])
(DOCS/'Thiet-ke-CSDL-Sprint3-4-buoc.md').write_text('\n'.join(M),encoding='utf-8')

# Word output: portrait text; landscape figure pages for legible ERDs.
d=Document();sec=d.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
for n in ['Normal','Body Text']:
 d.styles[n].font.name='Times New Roman';d.styles[n].font.size=Pt(12)
 d.styles[n].paragraph_format.space_after=Pt(6);d.styles[n].paragraph_format.line_spacing=1.1
for n,size in [('Heading 1',15),('Heading 2',13)]:
 d.styles[n].font.name='Times New Roman';d.styles[n].font.size=Pt(size);d.styles[n].font.color.rgb=RGBColor.from_string('17365D')
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('Lê Hoàng Lâm · S3-L1 | Trang ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def landscape(on):
 s=d.add_section(WD_SECTION_START.NEW_PAGE)
 s.orientation=WD_ORIENT.LANDSCAPE if on else WD_ORIENT.PORTRAIT
 s.page_width=Cm(29.7 if on else 21);s.page_height=Cm(21 if on else 29.7)
 s.left_margin=s.right_margin=Cm(2);s.top_margin=s.bottom_margin=Cm(1.5 if on else 2)
 if not on:s.left_margin=Cm(2.5)
for idx,e in enumerate(B):
 if e[0]=='h':
  if idx+1<len(B) and B[idx+1][0]=='image':continue
  d.add_heading(e[2],e[1])
 elif e[0]=='p':d.add_paragraph(e[1])
 elif e[0]=='code':
  pa=d.add_paragraph();r=pa.add_run(e[1]);r.font.name='Consolas';r.font.size=Pt(9)
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F2F2');pa._p.get_or_add_pPr().append(sh)
 elif e[0]=='table':
  headers,rows=e[1:];t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
  for c,x in zip(t.rows[0].cells,headers):c.text=x
  repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
  for row in rows:
   for c,x in zip(t.add_row().cells,row):c.text=str(x)
  for i,row in enumerate(t.rows):
   for cell in row.cells:
    for pa in cell.paragraphs:
     for r in pa.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=i==0
  d.add_paragraph()
 elif e[0]=='image':
  landscape(True)
  if idx>0 and B[idx-1][0]=='h':d.add_heading(B[idx-1][2],2)
  d.add_paragraph(e[2],'Caption')
  shape=d.add_picture(str(DOCS/e[1]),width=Cm(25))
  # Preserve full image; scale to maximum diagram viewport, never crop FK data.
  if shape.height>Cm(12):
   ratio=Cm(12)/shape.height;shape.width=int(shape.width*ratio);shape.height=Cm(12)
  d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
  for k,v in e[3]:
   pa=d.add_paragraph();pa.add_run(k+': ').bold=True;pa.add_run(v)
   for r in pa.runs:r.font.size=Pt(10)
  landscape(False)
d.core_properties.title='S3-L1 – Thiết kế CSDL Sprint 3 theo 4 bước'
d.core_properties.author='Lê Hoàng Lâm – có OpenAI Codex hỗ trợ'
d.core_properties.subject='Tài khoản, giỏ hàng, đặt hàng; ERD nối quản lý sách; thiết kế đề xuất'
d.save(OUT/'S3-L1 - Thiết kế CSDL Sprint 3 theo 4 bước.docx')
print('Generated Word, Markdown, JSON and two ERDs (PNG/SVG/DOT).')
